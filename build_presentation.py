"""
=============================================================================
IssueTrace: Comprehensive User Guide & System Walkthrough Presentation Builder
=============================================================================
This script generates an executive, high-fidelity 30-slide presentation deck
in widescreen (16:9) format explaining the IssueTrace Operational Workflow &
Reconciliation Platform for users (analysts, technicians, leads, managers).

Adheres to Antigravity Rule 15:
- All high-resolution screenshots are consolidated in `presentation_assets/`.
- Re-executable in a single command: `python build_presentation.py`.
=============================================================================
"""

import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # Blank slide

    base_dir = os.path.dirname(os.path.abspath(__file__))
    assets_dir = os.path.join(base_dir, "presentation_assets")

    # Brand & Palette Tokens
    BG_DARK = RGBColor(15, 23, 42)          # Slate 900
    BG_CARD_DARK = RGBColor(30, 41, 59)     # Slate 800
    BG_LIGHT = RGBColor(248, 250, 252)      # Slate 50
    CARD_BG_LIGHT = RGBColor(255, 255, 255) # Pure White
    TEXT_LIGHT = RGBColor(241, 245, 249)    # Slate 100
    TEXT_MUTED_DARK = RGBColor(148, 163, 184)# Slate 400
    TEXT_DARK = RGBColor(15, 23, 42)        # Slate 900
    TEXT_BODY = RGBColor(51, 65, 85)        # Slate 700
    TEXT_MUTED = RGBColor(100, 116, 139)    # Slate 500
    
    BRAND_BLUE = RGBColor(21, 93, 252)      # Electric Blue #155DFC
    BRAND_ORANGE = RGBColor(242, 116, 33)   # IssueTrace Orange #F27421
    ACCENT_EMERALD = RGBColor(16, 185, 129) # Emerald #10B981
    ACCENT_PURPLE = RGBColor(147, 51, 234)  # Purple #9333EA
    ACCENT_AMBER = RGBColor(217, 119, 6)    # Amber #D97706
    BORDER_LIGHT = RGBColor(226, 232, 240)  # Slate 200

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category, subtitle=None, dark_mode=False):
        # Category Pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(2.8), Inches(0.28))
        pill.fill.solid()
        pill.fill.fore_color.rgb = RGBColor(30, 58, 138) if dark_mode else RGBColor(239, 246, 255)
        pill.line.color.rgb = BRAND_BLUE
        pill.line.width = Pt(1)
        tf = pill.text_frame
        tf.word_wrap = False
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = category.upper()
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(147, 197, 253) if dark_mode else BRAND_BLUE

        # Slide Title
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.55))
        tf2 = tb.text_frame
        tf2.word_wrap = True
        tf2.margin_top = tf2.margin_bottom = tf2.margin_left = tf2.margin_right = 0
        p2 = tf2.paragraphs[0]
        run2 = p2.add_run()
        run2.text = title
        run2.font.size = Pt(20)
        run2.font.bold = True
        run2.font.color.rgb = TEXT_LIGHT if dark_mode else TEXT_DARK

        # Subtitle
        if subtitle:
            p_sub = tf2.add_paragraph()
            p_sub.space_before = Pt(2)
            r_sub = p_sub.add_run()
            r_sub.text = subtitle
            r_sub.font.size = Pt(10)
            r_sub.font.color.rgb = TEXT_MUTED_DARK if dark_mode else TEXT_MUTED

    def build_feature_slide(title, category, img_filename, summary_text, takeaways, user_tip=None, subtitle=None):
        slide = prs.slides.add_slide(blank_layout)
        set_slide_background(slide, BG_LIGHT)
        add_header(slide, title, category, subtitle=subtitle, dark_mode=False)

        # -------------------------------------------------------------
        # Left Column: Structured Feature Narrative & Actions
        # -------------------------------------------------------------
        left_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(4.35), Inches(5.8))
        ltf = left_box.text_frame
        ltf.word_wrap = True
        ltf.margin_top = ltf.margin_left = ltf.margin_right = ltf.margin_bottom = 0

        # Summary paragraph
        sp = ltf.paragraphs[0]
        s_run = sp.add_run()
        s_run.text = summary_text + "\n"
        s_run.font.size = Pt(11)
        s_run.font.color.rgb = TEXT_BODY

        # Feature takeaways
        for heading, bullet in takeaways:
            hp = ltf.add_paragraph()
            hp.space_before = Pt(8)
            hrun = hp.add_run()
            hrun.text = f"•  {heading}\n"
            hrun.font.size = Pt(10.5)
            hrun.font.bold = True
            hrun.font.color.rgb = BRAND_BLUE

            bp = ltf.add_paragraph()
            bp.space_before = Pt(1)
            brun = bp.add_run()
            brun.text = f"   {bullet}"
            brun.font.size = Pt(9.5)
            brun.font.color.rgb = TEXT_BODY

        # User action tip card at the bottom of left column
        if user_tip:
            tip_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.25), Inches(4.35), Inches(0.85))
            tip_box.fill.solid()
            tip_box.fill.fore_color.rgb = RGBColor(241, 245, 249) # Slate 100
            tip_box.line.color.rgb = RGBColor(203, 213, 225)      # Slate 300
            tip_box.line.width = Pt(1)
            ttf = tip_box.text_frame
            ttf.word_wrap = True
            ttf.margin_top = ttf.margin_bottom = Inches(0.08)
            ttf.margin_left = ttf.margin_right = Inches(0.12)
            tp1 = ttf.paragraphs[0]
            tr1 = tp1.add_run()
            tr1.text = "💡 USER ACTION GUIDE: "
            tr1.font.size = Pt(8.5)
            tr1.font.bold = True
            tr1.font.color.rgb = BRAND_ORANGE
            tr2 = tp1.add_run()
            tr2.text = user_tip
            tr2.font.size = Pt(8.5)
            tr2.font.color.rgb = TEXT_BODY

        # -------------------------------------------------------------
        # Right Column: Proportional High-Res Screenshot in Frame
        # -------------------------------------------------------------
        img_path = os.path.join(assets_dir, img_filename)
        frame_x = Inches(5.35)
        frame_y = Inches(1.35)
        max_w = Inches(7.18)
        max_h = Inches(5.75)

        # Outer card container
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, frame_x, frame_y, max_w, max_h)
        frame.fill.solid()
        frame.fill.fore_color.rgb = CARD_BG_LIGHT
        frame.line.color.rgb = BORDER_LIGHT
        frame.line.width = Pt(1.5)

        if os.path.exists(img_path):
            try:
                with Image.open(img_path) as im:
                    orig_w, orig_h = im.size
                aspect = orig_w / orig_h

                # Calculate fitted dimensions maintaining aspect ratio
                target_w = max_w - Inches(0.12)
                target_h = target_w / aspect

                if target_h > (max_h - Inches(0.12)):
                    target_h = max_h - Inches(0.12)
                    target_w = target_h * aspect

                # Center within container
                img_x = frame_x + (max_w - target_w) / 2
                img_y = frame_y + (max_h - target_h) / 2

                slide.shapes.add_picture(img_path, img_x, img_y, width=target_w, height=target_h)
            except Exception as e:
                print(f"Warning: could not render image {img_filename}: {e}")

    # =========================================================================
    # SLIDE 1: Hero Title Slide (Dark Theme)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, BG_DARK)

    # Accent Top Header Bar
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = BRAND_BLUE
    top_bar.line.fill.background()

    # Title & Subtitle
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.3), Inches(3.4))
    tf1 = title_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    r0 = p0.add_run()
    r0.text = "ENTERPRISE OPERATIONAL WORKFLOW & RECONCILIATION PLATFORM\n"
    r0.font.size = Pt(12)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(56, 189, 248) # Sky 400

    p1 = tf1.add_paragraph()
    p1.space_before = Pt(4)
    r1 = p1.add_run()
    r1.text = "IssueTrace: User Guide & System Walkthrough"
    r1.font.size = Pt(34)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf1.add_paragraph()
    p2.space_before = Pt(12)
    r2 = p2.add_run()
    r2.text = "An end-to-end visual manual for back-office settlement analysts, investigation technicians, squad leads, and compliance authorities. Designed for automated multi-source reconciliation, zero-hardcoded schema mappings, forensic transaction tracking, and strict Maker-Checker dual authorization."
    r2.font.size = Pt(13.5)
    r2.font.color.rgb = TEXT_MUTED_DARK

    # Pillars
    pillars = [
        ("Automated Ingestion", "Ingest multi-format files (Excel, CSV, XML), SFTP schedules, and direct partner intake queues."),
        ("Visual DAG Pipelines", "Assemble modular validation rules visually with instant pre-flight topology checks."),
        ("2-Level Investigation", "Forensic transaction trajectory tracing paired with deep cross-system database lookup."),
        ("Maker-Checker Dual Control", "Enforce strict anti-self-approval on all financial resolutions and discrepancy overrides.")
    ]
    card_w = Inches(2.68)
    card_gap = Inches(0.24)
    start_x = Inches(1.0)
    card_y = Inches(5.2)

    for i, (head, desc) in enumerate(pillars):
        c_x = start_x + i * (card_w + card_gap)
        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_x, card_y, card_w, Inches(1.65))
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD_DARK
        card.line.color.rgb = RGBColor(51, 65, 85)
        card.line.width = Pt(1)

        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_top = ctf.margin_left = ctf.margin_right = ctf.margin_bottom = Inches(0.14)
        cp1 = ctf.paragraphs[0]
        cr1 = cp1.add_run()
        cr1.text = head + "\n"
        cr1.font.size = Pt(12)
        cr1.font.bold = True
        cr1.font.color.rgb = RGBColor(255, 255, 255)

        cp2 = ctf.add_paragraph()
        cp2.space_before = Pt(3)
        cr2 = cp2.add_run()
        cr2.text = desc
        cr2.font.size = Pt(9.5)
        cr2.font.color.rgb = TEXT_MUTED_DARK

    # =========================================================================
    # SLIDE 2: End-to-End Operational Lifecycle
    # =========================================================================
    build_feature_slide(
        title="End-to-End Operational Lifecycle & Flow Architecture",
        category="Architecture & Flow",
        subtitle="Complete journey from raw feed staging to permanent reconciliation and sign-off",
        img_filename="44_end_to_end_flow.png",
        summary_text="IssueTrace orchestrates financial settlement through five deterministic stages, eliminating manual data handling and isolating data discrepancies.",
        takeaways=[
            ("Stage 1: Multi-Channel Ingress", "Extracts settlement data from SFTP, batch Excel/CSV, webhooks, or email intake queues into staging."),
            ("Stage 2: Schema Normalization", "Maps external columns to canonical directory fields and normalizes status values via dictionaries."),
            ("Stage 3: DAG Pipeline Validation", "Executes multi-step rule checks (existence, matching amounts, duplicate detection) in isolated unlogged mirrors."),
            ("Stage 4: Discrepancy Investigation", "Flags mismatched records into technician queues for forensic root-cause analysis and compare-and-resolve."),
            ("Stage 5: Maker-Checker Governance", "Routes proposed overrides through dual authorization before committing final resolutions.")
        ],
        user_tip="Follow this lifecycle sequentially when onboarding new settlement partners or troubleshooting transaction discrepancies."
    )

    # =========================================================================
    # SLIDE 3: Modern Design System & User Interface
    # =========================================================================
    build_feature_slide(
        title="Design System, Visual Tokens & Workspace Ergonomics",
        category="Design System",
        subtitle="High-contrast, responsive visual foundation built for operational clarity",
        img_filename="00_design_system.png",
        summary_text="The IssueTrace interface utilizes curated modern tokens, distinct color semantics, and stationary headers to provide an effortless user experience.",
        takeaways=[
            ("Semantic Color Coding", "Standardized badge colors across the platform: Emerald (Pass/Matched), Rose (Critical/Failed), Amber (High/Under Review), Purple (Escalated)."),
            ("Typography & Contrast Hierarchy", "Inter and Outfit fonts paired with high-contrast text ensure zero eye-strain during prolonged operational analysis."),
            ("Scroll-Contained Viewports", "Sticky table headers and docked pagination footers remain fixed on-screen during scrolling."),
            ("Adaptive Themes", "Seamlessly switches between high-visibility Light Mode and eye-friendly Dark Mode with instant persistence.")
        ],
        user_tip="Use the theme switcher in the bottom navigation to adapt the display to your workspace lighting environment."
    )

    # =========================================================================
    # SLIDE 4: Intake Cockpit & Multi-Channel Ingress
    # =========================================================================
    build_feature_slide(
        title="Intake Cockpit: External Request Ingress & Triage",
        category="Request Management",
        subtitle="Centralized inbox capturing inbound requests across email, webhooks, and partner channels",
        img_filename="19_intake.png",
        summary_text="The Intake Cockpit captures external settlement requests, dispute notices, and batch files, applying automated AI triage before task creation.",
        takeaways=[
            ("Multi-Channel Consolidation", "Unifies inbound messages from Microsoft 365, Zoho Mail, SFTP notifications, and banking webhooks."),
            ("AI Categorization & Urgency", "Automatically classifies requests (Reconciliation Exception, Chargeback, Fee Discrepancy) with confidence scores."),
            ("Interactive Paginated Queue", "Docked pagination controls (7 items per page default) with fast keyword search and status filters."),
            ("1-Click Task Conversion", "Operators can review raw payloads, inspect file attachments, and convert staged requests directly into operational cases.")
        ],
        user_tip="Filter by 'Ready for task' and sort by Urgency = 'Critical' at the start of your shift to triage urgent settlement exceptions first."
    )

    # =========================================================================
    # SLIDE 5: Automated SFTP File Staging
    # =========================================================================
    build_feature_slide(
        title="Remote SFTP File Staging & Ingestion Automation",
        category="Data Ingestion",
        subtitle="Scheduled extraction, rotating date navigation, and multi-format sheet parsing",
        img_filename="16_file_staging.png",
        summary_text="Automates the retrieval and ingestion of banking settlement files from remote SFTP servers, eliminating manual file downloads.",
        takeaways=[
            ("Dynamic Rotating Date Folders", "Auto-discovers calendar subfolders (e.g., /2026/sep/...) using dynamic wildcard matching."),
            ("Multi-Format Spreadsheet Parsing", "Seamlessly parses multi-sheet Excel (.xlsx, .xls), CSV, XML, and fixed-width text files."),
            ("Header & Offset Detection", "Accommodates merged report titles, banner metadata, and variable starting header rows effortlessly."),
            ("Unlogged Mirror Tables", "Loads raw data directly into typed PostgreSQL unlogged mirror tables for ultra-fast processing.")
        ],
        user_tip="Configure automated polling schedules on the SFTP configuration card so morning clearing files are staged before your squad arrives."
    )

    # =========================================================================
    # SLIDE 6: Global Data Directory & Column Mapping
    # =========================================================================
    build_feature_slide(
        title="Global Data Directory & Physical Column Mapping",
        category="Data Setup",
        subtitle="Agnostic dictionary architecture decoupling database schemas from workflow rules",
        img_filename="14_data_mapping.png",
        summary_text="The Global Standard Directory provides a canonical dictionary layer that shields business rules from physical column naming differences.",
        takeaways=[
            ("Zero Hardcoded Field Names", "Physical database columns (e.g., `TXN_AMT`, `AMOUNT_CENTS`) map dynamically to canonical keys (`amount`)."),
            ("Cross-System Portability", "Workflows referencing canonical keys run seamlessly across diverse bank core databases without rewriting SQL."),
            ("Visual Mapping Workbench", "Interactive drag-and-drop or dropdown column mapping with immediate validation of data types."),
            ("Database Scoping & Isolation", "Maintains distinct column configurations per external connection, preventing schema collisions.")
        ],
        user_tip="Always map required reconciliation columns (Transaction ID, Amount, Timestamp, Status) before activating a new database connection."
    )

    # =========================================================================
    # SLIDE 7: Value Dictionaries & Status Harmonization
    # =========================================================================
    build_feature_slide(
        title="Value Dictionary Workbench & Code Harmonization",
        category="Data Harmonization",
        subtitle="Standardizing heterogeneous response codes into uniform operational statuses",
        img_filename="37_new_value_dictionary.png",
        summary_text="Normalizes disparate status codes from payment switches, core banking engines, and partner networks into standardized business statuses.",
        takeaways=[
            ("Canonical Term Normalization", "Translates diverse codes (`00`, `APPR`, `APPROVED`, `PAID`) into a single canonical `SUCCESS` status."),
            ("Failure & Decline Mapping", "Maps switch decline codes (`51` = Insufficient Funds, `54` = Expired Card) to uniform investigation categories."),
            ("Multi-Dictionary Support", "Create specialized dictionary profiles for Card Networks, ACH Clearing, and Core Banking feeds."),
            ("Instant Rule Integration", "Validation boxes utilize dictionary lookups directly, simplifying rule definitions.")
        ],
        user_tip="Define default fallback values in your value dictionaries to capture unrecognized codes as 'UNKNOWN_EXCEPTION' for review."
    )

    # =========================================================================
    # SLIDE 8: Data Transformations & Calculated Expression Studio
    # =========================================================================
    build_feature_slide(
        title="Data Transformations & Calculated Expression Studio",
        category="Data Preparation",
        subtitle="On-the-fly arithmetic scaling, string formatting, and SQL expression transformations",
        img_filename="15_transformations.png",
        summary_text="Allows operators to apply mathematical, string, and temporal transformations to raw input data before reconciliation rules execute.",
        takeaways=[
            ("Currency & Scale Conversions", "Converts integer minor currency units (cents, cents-to-dollars) into decimal representations."),
            ("String Cleansing & Trimming", "Strips special characters, trailing spaces, and standardizes account formatting for clean matching."),
            ("Dynamic SQL Expression Builder", "Craft custom expressions (e.g., `CASE WHEN type = 'CR' THEN amount * -1 ELSE amount END`)."),
            ("Pre-Execution Preview", "Tests transformation formulas against sample records to verify calculation accuracy before saving.")
        ],
        user_tip="Apply amount scaling transformations in the preparation step to ensure input currencies align exactly with core banking ledgers."
    )

    # =========================================================================
    # SLIDE 9: Operations Workspace & Discrepancy Queue
    # =========================================================================
    build_feature_slide(
        title="Operations Workspace & High-Throughput Case Grid",
        category="Daily Operations",
        subtitle="Mission control for settlement analysts to track, filter, and inspect transaction batches",
        img_filename="01_cases_workspace.png",
        summary_text="The primary working environment for analysts, providing real-time transaction ingestion, status indicators, and batch action controls.",
        takeaways=[
            ("Real-Time Discrepancy Queue", "Ingests batch records and displays validation badges, error classifications, and match states."),
            ("Customizable Column Visibility", "Toggle financial, account, fee, and timestamp columns with persistent user preference storage."),
            ("Multi-Dimensional Filtering", "Filter instantly by status (New, Processing, Discrepancy), currency, partner institution, or date range."),
            ("One-Click Batch Execution", "Trigger validation runs, force re-runs on staged items, revert tasks, or export audit CSV reports with a single click.")
        ],
        user_tip="Customize your visible columns to keep Transaction ID, Amount, Status, and Error Verdict in view for quick scanning."
    )

    # =========================================================================
    # SLIDE 10: Case Staging & Task Creation Popup
    # =========================================================================
    build_feature_slide(
        title="Staging New Cases & Pipeline Assignment",
        category="Case Creation",
        subtitle="Uploading raw settlement batches, attaching validation pipelines, and initiating runs",
        img_filename="21_new_case_upload_clean.png",
        summary_text="Provides an intuitive modal interface to upload new reconciliation files, assign execution pipelines, and configure batch metadata.",
        takeaways=[
            ("Drag-and-Drop Batch Upload", "Drop raw Excel, CSV, or TXT settlement files directly into the upload dropzone."),
            ("Pipeline & Workflow Binding", "Select from available visual DAG workflows (e.g., Daily Settlement, Card Interchange Reconciliation)."),
            ("Database Scope Selection", "Explicitly choose the target database connection and core ledger table to reconcile against."),
            ("Execution Policy Configuration", "Set batch priorities (Urgent, Standard) and choose whether to auto-run validation immediately.")
        ],
        user_tip="Verify the target database connection before submitting the staging modal to ensure reconciliation matches the correct core ledger."
    )

    # =========================================================================
    # SLIDE 11: Real-Time Pipeline Execution & Progress Monitoring
    # =========================================================================
    build_feature_slide(
        title="Real-Time Pipeline Execution & Progress Monitoring",
        category="Execution Engine",
        subtitle="Live telemetry, stage progress milestones, and non-blocking asynchronous processing",
        img_filename="22_investigation_running.png",
        summary_text="Tracks active validation execution in real time with visual progress milestones, stage transitions, and execution telemetry.",
        takeaways=[
            ("Non-Blocking Background Processing", "Executes multi-thousand transaction batches asynchronously without freezing the user interface."),
            ("Stage-by-Stage Milestone Bars", "Displays progress across Ingestion, Rule Evaluation, Mirror Reconciliation, and Output Reporting."),
            ("Live Counter Telemetry", "Monitors processed records, passing transactions, detected discrepancies, and syntax errors in real time."),
            ("Safe Task Cancellation", "Operators can pause or terminate running batches if an incorrect file was uploaded.")
        ],
        user_tip="You can navigate freely across other workspaces while a batch executes; progress will continue and notify you upon completion."
    )

    # =========================================================================
    # SLIDE 12: Granular Transaction Forensics Drawer
    # =========================================================================
    build_feature_slide(
        title="Transaction Forensics & Sequential Rule Trajectory",
        category="Forensic Investigation",
        subtitle="Row-level deep inspection detailing input data, mirror records, and rule evaluation verdicts",
        img_filename="02_transaction_drawer.png",
        summary_text="Opening any transaction launches the forensic drawer, revealing the exact sequence of rules evaluated and their outcomes.",
        takeaways=[
            ("Sequential Trajectory Audit", "Traces step-by-step verdicts (PASS, FAIL, ERROR) mapped directly to pipeline actions (CONTINUE, STOP, FLAG)."),
            ("Source vs. Mirror Payload Inspection", "Compare raw ingested fields side-by-side with database mirror records down to the microsecond."),
            ("Execution Latency & Metadata", "Inspect execution duration, evaluated SQL criteria, and parameter values for each validation step."),
            ("Direct Action Triggers", "Launch deep compare-and-resolve, initiate team discussions, or stage resolution proposals directly from the drawer.")
        ],
        user_tip="Click on any failed step in the trajectory list to view the exact SQL criteria and parameter bounds evaluated by the engine."
    )

    # =========================================================================
    # SLIDE 13: Two-Level Forensic Investigation Workflow
    # =========================================================================
    build_feature_slide(
        title="Two-Level Discrepancy Investigation Workflow",
        category="Advanced Analysis",
        subtitle="Hierarchical forensic analysis: batch coordinate checks vs. deep core ledger lookups",
        img_filename="39_case_two_levels_investigate.png",
        summary_text="Separates investigation into Level 1 (batch coordinate validation) and Level 2 (deep core banking ledger reconciliation).",
        takeaways=[
            ("Level 1: Batch Coordinate Audit", "Validates row coordinate integrity, batch sum totals, file headers, and internal transaction references."),
            ("Level 2: Deep Core Banking Lookup", "Executes batched tuple queries against external core banking ledgers to detect missing settlements."),
            ("Clear Investigation Statuses", "Tracks investigation progress: Uninvestigated, In Progress, Discrepancy Confirmed, Resolved."),
            ("Root Cause Categorization", "Classify root causes: Timing Difference, Network Dropout, Fee Miscalculation, or Legitimate Fraud.")
        ],
        user_tip="Complete Level 1 checks first; if internal coordinates match, escalate to Level 2 deep lookup to identify external ledger mismatches."
    )

    # =========================================================================
    # SLIDE 14: Compare & Resolve Workbench
    # =========================================================================
    build_feature_slide(
        title="Compare & Resolve: Side-by-Side Discrepancy Reconciliation",
        category="Discrepancy Resolution",
        subtitle="Visual field-level comparison, discrepancy highlighting, and resolution proposal staging",
        img_filename="23_compare_and_resolve.png",
        summary_text="The Compare & Resolve workbench gives technicians a visual diff between the source record and core ledger, enabling manual resolution proposals.",
        takeaways=[
            ("Field-by-Field Visual Diff", "Highlights mismatched fields in bold amber/red (e.g., Ingested $1,500.00 vs. Core Ledger $1,450.00)."),
            ("Automated Resolution Options", "Select from standardized operational resolutions: Force Match, Write-Off, Manual Reversal, or Reject."),
            ("Mandatory Justification Notes", "Technicians must provide detailed rationale and operational notes before submitting a proposal."),
            ("Snapshot Locking", "Locks transaction state in `PENDING_CHECKER_REVIEW` with an immutable evidence snapshot for auditability.")
        ],
        user_tip="Always attach the core banking reference number in the resolution notes to assist the Checker during dual-control review."
    )

    # =========================================================================
    # SLIDE 15: Workflow Studio & Visual DAG Canvas
    # =========================================================================
    build_feature_slide(
        title="Workflow Studio: Visual Flowchart DAG Pipeline Builder",
        category="Pipeline Engineering",
        subtitle="No-code visual canvas to design, wire, and persist multi-stage reconciliation pipelines",
        img_filename="03_workflow_studio.png",
        summary_text="Workflow Studio allows administrators and engineers to construct sophisticated reconciliation pipelines visually without writing code.",
        takeaways=[
            ("Interactive Flowchart Canvas", "Drag and drop Start, End, and modular validation nodes; connect ports with visual SVG bezier curves."),
            ("Instant Pre-Flight Topology Validator", "Validates graph topology continuously, alerting on disconnected nodes, orphan branches, or infinite loops."),
            ("Database & Stage Scoping", "Assign target databases and cardinality policies (Single Row, Multi-Row Aggregate) per stage."),
            ("PostgreSQL Relational Persistence", "Stores pipeline definitions directly in PostgreSQL tables for reproducible, version-controlled execution.")
        ],
        user_tip="Click 'Validate Graph' in the canvas toolbar before saving to ensure all nodes and branch conditions are wired properly."
    )

    # =========================================================================
    # SLIDE 16: Modular Validation Rules Library & Sandbox
    # =========================================================================
    build_feature_slide(
        title="Modular Validation Rules Library & Sandbox",
        category="Rule Architecture",
        subtitle="Reusable validation building blocks with built-in interactive unit testing",
        img_filename="04_rules_library.png",
        summary_text="A centralized repository of tested validation building blocks that can be dragged into any workflow canvas across the organization.",
        takeaways=[
            ("Four Canonical Box Types", "Strictly structured into Ingestion Search, Condition Check, Reconciliation, and Report Output blocks."),
            ("Standardized Rule Templates", "Pre-built templates for ISO Decline Codes, Existence Checks, Amount Thresholds, and Date Consistency."),
            ("Interactive Sandbox Testing", "Run dry-run unit tests directly from the library cards against live sample rows before publishing."),
            ("Universal Hashtag Tagging", "Tag validation boxes with hashtags (e.g., #card-reconciliation) for instant discovery across squads.")
        ],
        user_tip="Test any newly configured validation rule in the dry-run sandbox before adding it to an active production pipeline."
    )

    # =========================================================================
    # SLIDE 17: Multi-Leg Transaction Rules & Duplicate Elimination
    # =========================================================================
    build_feature_slide(
        title="Multi-Leg Transaction Rules & Duplicate Elimination",
        category="Integrity Rules",
        subtitle="Advanced partition grouping keys and conditional row criteria to eliminate false duplicates",
        img_filename="05_duplicate_check_builder.png",
        summary_text="Solves the challenge of complex multi-leg payments by grouping transactions by partition keys and filtering occurrences conditionally.",
        takeaways=[
            ("Common Partition Grouping Keys", "Groups transactions by correlation keys (e.g., `TR_GROUP_ID`, `SESSION_ID`) rather than flat row evaluation."),
            ("Conditional Leg Criteria Builder", "Specifies distinct criteria for specific transaction legs (e.g., `BO_TYPE = 'CMTP254'`) using standard operators."),
            ("Filtered Occurrence Counting", "Counts only the rows satisfying leg criteria within each partition group and evaluates thresholds (`COUNT = 1`)."),
            ("Multi-Leg Integrity Guarantee", "Eliminates false-positive duplicate alerts on legitimate multi-leg transfers and currency exchanges.")
        ],
        user_tip="When configuring duplicate checks on multi-leg transfers, specify the unique transaction group identifier as the partition key."
    )

    # =========================================================================
    # SLIDE 18: Automations: Event Triggers & Scheduled Batch Jobs
    # =========================================================================
    build_feature_slide(
        title="Automations: Event Triggers & Scheduled Batch Jobs",
        category="Process Automation",
        subtitle="Configuring cron schedules, event triggers, automated notifications, and auto-triage",
        img_filename="06_automations.png",
        summary_text="Automates recurring operational tasks, file polling, and reconciliation runs with configurable cron expressions and event listeners.",
        takeaways=[
            ("Flexible Trigger Types", "Trigger workflows based on cron schedules (e.g., every morning at 06:00 AM), file arrivals, or webhook events."),
            ("Automated Squad Routing", "Auto-assigns newly generated cases to specific squads (Cards, Settlement, Dispute Desk) based on criteria."),
            ("SLA Breach Alerting", "Triggers high-priority notifications when aging cases approach SLA expiration thresholds."),
            ("Execution History & Audit Log", "Maintains complete run history with success/failure statuses, error traces, and execution durations.")
        ],
        user_tip="Enable email or Teams notifications on mission-critical automation jobs so squad leads are notified of any polling interruptions."
    )

    # =========================================================================
    # SLIDE 19: Query Sandbox & Database Explorer
    # =========================================================================
    build_feature_slide(
        title="Query Sandbox & Database Explorer",
        category="Technical Analysis",
        subtitle="Safe read-only SQL workbench for ad-hoc transaction queries and data schema inspection",
        img_filename="26_query_sandbox.png",
        summary_text="Empowers technicians and database administrators to inspect database schemas and run ad-hoc read-only queries safely.",
        takeaways=[
            ("Safe Read-Only Guardrails", "Strictly enforces `SELECT` operations; mutating statements (UPDATE, DELETE, DROP) are blocked automatically."),
            ("Multi-Database Connection Switcher", "Query across Core Banking, Switch Archives, and Staged Mirror tables from a single workbench."),
            ("SQL Syntax Highlighting", "Modern editor with auto-completion, keyword coloring, line numbering, and query formatting."),
            ("Projected Results & CSV Export", "Instant tabular data grid with execution timing, row counts, and one-click export for investigation evidence.")
        ],
        user_tip="Use the Query Sandbox to verify customer account balances or check historical reference records during deep forensic investigations."
    )

    # =========================================================================
    # SLIDE 20: Team Tasks: Operational Kanban & Task Tracking
    # =========================================================================
    build_feature_slide(
        title="Team Tasks: Operational Kanban & Task Tracking",
        category="Squad Collaboration",
        subtitle="Visual workflow board consolidating squad assignments, priorities, and SLA timers",
        img_filename="08_team_tasks.png",
        summary_text="Provides an operational Kanban board tracking cases across Backlog, In Progress, Peer Review, and Completed columns.",
        takeaways=[
            ("Visual Kanban Columns", "Move cases across workflow states with clear priority indicators and assignee avatars."),
            ("Public Team Task Consolidation", "Consolidates individual operator tasks into team-visible boards for balanced workload distribution."),
            ("Real-Time SLA Timers", "Color-coded countdown badges alert squads to pending SLA breaches before deadlines are missed."),
            ("1-Click Workspace Navigation", "Clicking any task card immediately navigates directly to that case in the operations workspace.")
        ],
        user_tip="Review the Team Tasks board during daily standups to redistribute cases from overloaded operators to available squad members."
    )

    # =========================================================================
    # SLIDE 21: Shared Team Mailbox & Dispute Desk
    # =========================================================================
    build_feature_slide(
        title="Shared Team Mailbox & Customer Dispute Triage",
        category="Customer Operations",
        subtitle="Unified multi-pane email client connecting directly to team mailboxes via OAuth 2.0",
        img_filename="09_team_mailbox.png",
        summary_text="A dedicated dispute desk interface connecting to shared team mailboxes for handling inbound customer chargebacks and partner inquiries.",
        takeaways=[
            ("Multi-Pane Email Interface", "Thread list on the left, full message history and attachment preview pane on the right."),
            ("Centralized OAuth 2.0 Security", "Connects securely to Microsoft 365 or Google Workspace without exposing secrets to operators."),
            ("Attachment Ingestion", "Extract and preview attached customer proof-of-payment receipts, bank letters, and chargeback claims."),
            ("Direct Case Creation", "Convert email threads directly into operational reconciliation cases with full message history attached.")
        ],
        user_tip="Use the 'Convert to Case' button in the message header to transform a dispute email into an active investigation task."
    )

    # =========================================================================
    # SLIDE 22: Contextual Case Discussions & Slash Directives
    # =========================================================================
    build_feature_slide(
        title="Contextual Case Discussions & Slash Directives",
        category="In-App Chat",
        subtitle="In-context messaging drawer supporting user mentions, squad tags, hashtags, and slash commands",
        img_filename="27_case_discussion.png",
        summary_text="Each case contains an integrated discussion drawer where operators collaborate in context without leaving the transaction screen.",
        takeaways=[
            ("Contextual Discussion Drawer", "Chat history is permanently bound to the specific case ID for an immutable collaboration record."),
            ("Rich Mention Directives", "Use `@username` to mention colleagues, `@@team` to notify entire squads, and `#hashtag` to link related workflows."),
            ("Slash Command Execution", "Type `/create-escalate` to open an inter-team room or `/resolve` to initiate a resolution proposal."),
            ("Visual Badge Formatting", "Mentions and directives are styled with distinct, colorful pill badges for maximum readability.")
        ],
        user_tip="Tag the relevant squad using '@@team-cards' or '@@team-core-banking' when requesting cross-department transaction verification."
    )

    # =========================================================================
    # SLIDE 23: Inter-Team Escalation Rooms & Cross-Squad Collaboration
    # =========================================================================
    build_feature_slide(
        title="Inter-Team Escalation Rooms & Cross-Squad Collaboration",
        category="Escalation Management",
        subtitle="Dedicated shared rooms bridging operational squads to resolve complex discrepancies",
        img_filename="30_room_escalating_team.png",
        summary_text="When a discrepancy spans multiple departments (e.g., Settlement ↔ Core Banking), operators spawn dedicated Escalation Rooms.",
        takeaways=[
            ("Multi-Department Bridge", "Creates a shared workspace uniting the escalating squad and receiving squad under a single ticket."),
            ("Frozen Evidence Snapshots", "Both teams inspect the exact same immutable snapshot of transaction data and mirror records."),
            ("Real-Time Unified Chat", "Full directive autocompletion (`@`, `@@`, `#`, `/`) ensures rapid cross-team communication."),
            ("Clear Squad Ownership", "Explicit tracking of which squad currently holds the ball to prevent resolution delays.")
        ],
        user_tip="Create an escalation room whenever a transaction error originates in an external system outside your squad's direct jurisdiction."
    )

    # =========================================================================
    # SLIDE 24: Escalation Agreements & Mutual Resolution Commitments
    # =========================================================================
    build_feature_slide(
        title="Escalation Agreements & Mutual Resolution Commitments",
        category="Formal Settlement",
        subtitle="Formalizing resolution terms with dual-team sign-off before closing tickets",
        img_filename="32_escalation_agreement.png",
        summary_text="To prevent premature case closures, IssueTrace requires formal Escalation Agreements with sign-offs from both squads.",
        takeaways=[
            ("Formal Agreement Terms", "Documents agreed resolution steps (e.g., Ledger Adjustment, Chargeback Reversal, GL Write-Off)."),
            ("Dual-Team Authorization", "Requires verified approval from authorized leads of both squads before the case can close."),
            ("Immutable Audit Binding", "Agreement text, timestamps, and approver identities are cryptographically recorded in the case history."),
            ("Automated Task Unblocking", "Once both leads approve, the parent case automatically advances to resolved status.")
        ],
        user_tip="Ensure the agreement terms clearly specify which team is executing the physical general ledger entry in the core banking system."
    )

    # =========================================================================
    # SLIDE 25: Technician Solution Scripts & Universal Hashtags
    # =========================================================================
    build_feature_slide(
        title="Technician Solution Scripts & Universal Hashtag Knowledge",
        category="Knowledge Management",
        subtitle="Publishing and discovering reusable SQL remediation scripts via universal hashtags",
        img_filename="40_solution_script.png",
        summary_text="Allows experienced technicians to capture SQL remediation scripts and publish them under universal hashtags for squad reuse.",
        takeaways=[
            ("Remediation Script Repository", "Save tested SQL scripts that correct recurring settlement anomalies (e.g., orphaned authorization records)."),
            ("Universal Hashtag Binding", "Tag scripts with hashtags (e.g., `#AIB_SETTLEMENT_2026`) linking workflows, DAG nodes, and discussions."),
            ("Squad Knowledge Sharing", "Junior analysts can discover and review vetted solutions authored by senior technicians."),
            ("Attribution & Governance", "Every script retains author attribution, revision timestamps, and safety execution notes.")
        ],
        user_tip="Search the universal hashtag bar before troubleshooting an unfamiliar error to see if a verified solution script already exists."
    )

    # =========================================================================
    # SLIDE 26: Maker-Checker Dual Authorization & Authority Center
    # =========================================================================
    build_feature_slide(
        title="Maker-Checker Dual Authorization & Authority Center",
        category="Governance & Compliance",
        subtitle="Strict Four-Eyes Principle enforcing anti-self-approval on all financial overrides",
        img_filename="18_approvals.png",
        summary_text="The Authority Center enforces banking-grade governance, requiring two distinct operators to authorize any financial adjustment.",
        takeaways=[
            ("Strict Anti-Self-Approval", "The maker who proposes an override cannot approve it (`makerId !== checkerId`); enforced at the API layer."),
            ("Pending Review Queue", "Supervisors review proposed overrides (Force Match, Write-Off, Manual Reversal) with evidence snapshots."),
            ("Frozen Evidence Snapshots", "Approvers inspect the exact transaction state, source values, and audit logs recorded when proposed."),
            ("Complete Audit Trail", "All approval and rejection decisions are permanently logged with timestamps, comments, and operator credentials.")
        ],
        user_tip="Supervisors should review the Approvals Inbox multiple times daily to prevent pending transaction batches from aging."
    )

    # =========================================================================
    # SLIDE 27: Notification Center & Multi-Channel Alert Matrix
    # =========================================================================
    build_feature_slide(
        title="Notification Center & Multi-Channel Alert Matrix",
        category="Notification System",
        subtitle="Real-time alert popover, dedicated full-page inbox, and granular preference controls",
        img_filename="46_notification_center.png",
        summary_text="A comprehensive notification system ensuring operators and leads never miss urgent SLA breaches, mentions, or approvals.",
        takeaways=[
            ("Top-Bar Popover & Full Center", "Quickly review notifications via the header bell popover or dive into the full tabbed notification center."),
            ("Filtered Categorization", "Organized tabs: All Notifications, Unread, Mentions, SLA Alerts, and System Announcements."),
            ("Granular Channel Preferences", "Customize delivery channels (In-App, Email, Push) independently for each notification category."),
            ("Immediate Deep Linking", "Clicking any notification navigates directly to the relevant case, discussion thread, or approval review.")
        ],
        user_tip="Set SLA breach alerts to both In-App and Email in your Notification Preferences to guarantee rapid response times."
    )

    # =========================================================================
    # SLIDE 28: Managerial Dashboard & Executive SLA Analytics
    # =========================================================================
    build_feature_slide(
        title="Managerial Dashboard & Executive SLA Analytics",
        category="Executive Reporting",
        subtitle="High-level operational metrics, team throughput velocity, and reconciliation trends",
        img_filename="11_dashboard.png",
        summary_text="Provides executives and operations heads with a centralized view of daily settlement volume, resolution velocity, and SLA adherence.",
        takeaways=[
            ("Executive KPI Summary", "Real-time metrics: Total Inbound Volume, Open Discrepancies, Resolved Transactions, and SLA Compliance Rate."),
            ("Weekly Throughput Trends", "Visual charts comparing incoming transaction volume against completed squad resolutions over time."),
            ("SLA Risk Monitoring", "Identifies aging cases at risk of breach, allowing management to reassign resources proactively."),
            ("Regulatory CSV Export", "Generate consolidated audit reports with a single click for financial regulators and internal compliance.")
        ],
        user_tip="Export the monthly summary report on the 1st of each month for executive leadership reviews and compliance filings."
    )

    # =========================================================================
    # SLIDE 29: Enterprise Administration, RBAC & Connection Security
    # =========================================================================
    build_feature_slide(
        title="Enterprise Administration, RBAC & Connection Security",
        category="System Administration",
        subtitle="Granular role privileges, team connection isolation, and centralized OAuth 2.0 management",
        img_filename="12_people_access.png",
        summary_text="Ensures multi-tenant security and compliance through strict role-based access control, team isolation, and credential masking.",
        takeaways=[
            ("Granular Role Permissions", "Assign distinct roles (Admin, Operations Analyst, Technician, Checker, Read-Only Auditor)."),
            ("Team Connection Isolation", "Database connections and mirror tables created within a squad are strictly isolated to authorized members."),
            ("Centralized OAuth 2.0 Management", "Enterprise credentials (Client ID, Secrets) are maintained centrally at Admin level with masked secrets."),
            ("Audit Logging & Health", "Live connectivity monitoring across database endpoints with real-time socket health and execution logging.")
        ],
        user_tip="Assign the 'Checker' role only to team leads and senior supervisors authorized to approve manual financial write-offs."
    )

    # =========================================================================
    # SLIDE 30: User Journey Cheat Sheet & Best Practices (Dark Theme)
    # =========================================================================
    slide30 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide30, BG_DARK)
    add_header(slide30, "Daily User Workflow & Reconciliation Best Practices", "Operational Summary", subtitle="Standard operating procedure for back-office teams", dark_mode=True)

    sop_steps = [
        ("1. Intake & Staging", "Review Intake queue for external requests; verify morning SFTP file staging completed successfully."),
        ("2. Pipeline Validation", "Execute visual DAG workflows against staged settlement files; monitor progress milestones in real time."),
        ("3. Discrepancy Review", "Filter Operations Workspace for failed transactions; inspect row-level trajectory forensics in the drawer."),
        ("4. Compare & Resolve", "Open two-level investigation; visually diff against core banking ledger; propose structured resolution."),
        ("5. Squad Collaboration", "Use in-case chat (`@user`, `@@team`, `/create-escalate`) or escalation rooms for cross-team alignment."),
        ("6. Checker Sign-Off", "Checker supervisor reviews proposed override in Authority Center; approves to commit final settlement.")
    ]

    card_w2 = Inches(5.6)
    card_h2 = Inches(1.55)
    gap_x2 = Inches(0.5)
    gap_y2 = Inches(0.22)
    s_x2 = Inches(0.8)
    s_y2 = Inches(1.5)

    for i, (sop_title, sop_desc) in enumerate(sop_steps):
        col = i % 2
        row = i // 2
        x = s_x2 + col * (card_w2 + gap_x2)
        y = s_y2 + row * (card_h2 + gap_y2)

        box = slide30.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, card_w2, card_h2)
        box.fill.solid()
        box.fill.fore_color.rgb = BG_CARD_DARK
        box.line.color.rgb = RGBColor(51, 65, 85)
        box.line.width = Pt(1)

        btf = box.text_frame
        btf.word_wrap = True
        btf.margin_top = btf.margin_left = btf.margin_right = btf.margin_bottom = Inches(0.14)

        bp1 = btf.paragraphs[0]
        br1 = bp1.add_run()
        br1.text = sop_title + "\n"
        br1.font.size = Pt(12)
        br1.font.bold = True
        br1.font.color.rgb = RGBColor(56, 189, 248) # Sky 400

        bp2 = btf.add_paragraph()
        bp2.space_before = Pt(3)
        br2 = bp2.add_run()
        br2.text = sop_desc
        br2.font.size = Pt(9.5)
        br2.font.color.rgb = TEXT_MUTED_DARK

    # Save presentation to repository root
    out_file1 = os.path.join(base_dir, "IssueTrace_User_Guide_and_System_Walkthrough.pptx")
    out_file2 = os.path.join(base_dir, "IssueTrace_System_Introduction.pptx")
    
    prs.save(out_file1)
    prs.save(out_file2)
    print(f"Presentation successfully generated and saved to:")
    print(f"  - {out_file1}")
    print(f"  - {out_file2}")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == "__main__":
    build_deck()
